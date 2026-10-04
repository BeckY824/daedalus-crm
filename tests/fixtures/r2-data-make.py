"""
r2-data 导入边界用的夹具。用 openpyxl 真写出来（再按需改几处 XML），不是自己拼 zip。
重新生成：python3 tests/fixtures/r2-data-make.py
"""
import datetime, io, os, zipfile
import openpyxl
from openpyxl.utils.datetime import CALENDAR_MAC_1904

D = os.path.dirname(os.path.abspath(__file__))
P = lambda n: os.path.join(D, n)


def save(wb, name):
    wb.save(P(name))


# 1. 表头在第二行：第一行是合并的大标题
wb = openpyxl.Workbook(); ws = wb.active
ws["A1"] = "2026 年 9 月客户名单"; ws.merge_cells("A1:D1")
ws.append(["姓名", "手机号", "公司", "备注"])
ws.append(["张三", "13800000001", "远山资本", "老客户"])
ws.append(["李四", "13800000002", "平川科技", ""])
save(wb, "r2-data-表头第二行.xlsx")

# 2. 合并单元格：同一家公司的三个人，公司那一格竖着合并
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["姓名", "手机号", "公司"])
ws.append(["张三", "13800000001", "远山资本"])
ws.append(["李四", "13800000002", None])
ws.append(["王五", "13800000003", None])
ws.merge_cells("C2:C4")
save(wb, "r2-data-合并单元格.xlsx")

# 3. 公式：手机号是公式。B2 带 Excel 缓存值（Excel 存过的样子），B3 没有缓存值（程序生成、没在 Excel 里打开过）
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["姓名", "手机号"])
ws.append(["张三", '="138"&"00000001"'])
ws.append(["李四", '="138"&"00000002"'])
buf = io.BytesIO(); wb.save(buf)
zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
zout = zipfile.ZipFile(P("r2-data-公式.xlsx"), "w", zipfile.ZIP_DEFLATED)
for it in zin.infolist():
    data = zin.read(it.filename)
    if it.filename == "xl/worksheets/sheet1.xml":
        s = data.decode()
        # 给 B2 补上 Excel 会写的缓存值（t="str" + <v>），B3 保持 openpyxl 原样（空的 <v>）
        s = s.replace('<c r="B2"><f>"138"&amp;"00000001"</f><v></v>', '<c r="B2" t="str"><f>"138"&amp;"00000001"</f><v>13800000001</v>', 1)
        data = s.encode()
    zout.writestr(it, data)
zout.close()

# 4. 1904 日期系统（老版 Mac Excel / Numbers 存出来的）
wb = openpyxl.Workbook(); wb.epoch = CALENDAR_MAC_1904; ws = wb.active
ws.append(["姓名", "手机号", "预计签约"])
ws.append(["张三", "13800000001", datetime.datetime(2026, 9, 19)])
save(wb, "r2-data-1904.xlsx")

# 5. 没对上的列里有日期：「加微信日期」这一列会并进备注
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["姓名", "手机号", "加微信日期"])
ws.append(["张三", "13800000001", datetime.datetime(2026, 9, 19)])
ws["C2"].number_format = "yyyy-mm-dd"
save(wb, "r2-data-日期并进备注.xlsx")

# 6. 中间一行是空的但设了行高：写出来是自闭合的 <row …/>
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["姓名", "手机号"])
ws.append(["张三", "13800000001"])
ws.row_dimensions[3].height = 30
ws.cell(row=4, column=1, value="李四"); ws.cell(row=4, column=2, value="13800000002")
ws.cell(row=5, column=1, value="王五"); ws.cell(row=5, column=2, value="13800000003")
buf = io.BytesIO(); wb.save(buf)
zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
zout = zipfile.ZipFile(P("r2-data-空行自闭合.xlsx"), "w", zipfile.ZIP_DEFLATED)
for it in zin.infolist():
    data = zin.read(it.filename)
    if it.filename == "xl/worksheets/sheet1.xml":
        # Excel 把「设了格式、没填内容」的行写成自闭合的 <row …/>
        data = data.decode().replace('<row r="3" ht="30" customHeight="1"></row>', '<row r="3" spans="1:2" ht="30" customHeight="1"/>').encode()
    zout.writestr(it, data)
zout.close()

# 7. 超宽：60 列，手机号在第 55 列
wb = openpyxl.Workbook(); ws = wb.active
head = [f"列{i}" for i in range(1, 61)]; head[0] = "姓名"; head[54] = "手机号"
ws.append(head)
row = [f"v{i}" for i in range(1, 61)]; row[0] = "张三"; row[54] = "13800000001"
ws.append(row)
save(wb, "r2-data-超宽.xlsx")

# 8. 上万行：10050 行数据
wb = openpyxl.Workbook(write_only=True); ws = wb.create_sheet()
ws.append(["姓名", "手机号"])
for i in range(10050):
    ws.append([f"客户{i}", f"137{i:08d}"])
save(wb, "r2-data-上万行.xlsx")

# 9. 只有表头
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["姓名", "手机号"])
save(wb, "r2-data-只有表头.xlsx")

# 10. 文本编码：GBK 的 csv、Excel「Unicode 文本」（UTF-16LE + BOM、制表符分隔）
text = "姓名,手机号,公司\n张三,13800000001,远山资本\n"
open(P("r2-data-gbk.csv"), "wb").write(text.encode("gbk"))
open(P("r2-data-utf16.txt"), "wb").write(b"\xff\xfe" + text.replace(",", "\t").replace("\n", "\r\n").encode("utf-16-le"))
open(P("r2-data-utf8bom.csv"), "wb").write(b"\xef\xbb\xbf" + text.encode("utf-8"))
# 11. 1904 日期系统里，数字存的号码、金额不能跟着日期一起挪（J-054）：只有日期格式的格子加 1462 天
wb = openpyxl.Workbook(); wb.epoch = CALENDAR_MAC_1904; ws = wb.active
ws.append(["姓名", "手机号", "预计签约", "加微信日期", "预算"])
ws.append(["张三", 13800000001, datetime.datetime(2026, 9, 19), datetime.datetime(2026, 9, 1), 50000])
ws["C2"].number_format = "yyyy-mm-dd"
ws["D2"].number_format = "mm-dd-yy"  # 内置格式 14
save(wb, "r2-data-1904-号码.xlsx")

# 12. 报错行号要和 Excel 里看到的一样（J-058）：第 1 行大标题、第 2 行表头、第 4 行空着
wb = openpyxl.Workbook(); ws = wb.active
ws["A1"] = "2026 年 9 月客户名单"; ws.merge_cells("A1:B1")
ws["A2"] = "姓名"; ws["B2"] = "手机号"
ws["A3"] = "张三"; ws["B3"] = "13800000001"
ws["A5"] = "李四"
ws["A6"] = "王五"; ws["B6"] = "123"
save(wb, "r2-data-行号.xlsx")

print("ok")
