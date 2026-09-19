#!/usr/bin/env python3
"""
重新生成 tests/fixtures/*.xlsx。

    pip install openpyxl && python3 scripts/make-xlsx-fixtures.py

**故意用 openpyxl 而不是我们自己的代码写。** 自己写自己读的测试只能证明
两段代码对同一个误解是一致的；用别人的写法生成，才验得到「真实的 Excel 文件
长什么样」——尤其是空格子在 XML 里直接不出现、以及共享字符串表这两件。
"""
from openpyxl import Workbook
from datetime import date
import os, zipfile

out = os.path.join(os.path.dirname(__file__), "..", "tests", "fixtures")
os.makedirs(out, exist_ok=True)

# 1) 一张普通的名单：中文、空格子、日期、被存成数字的手机号、要转义的引号
wb = Workbook()
ws = wb.active
ws.title = "名单"
ws.append(["姓名", "手机号", "公司", "预计签约", "备注"])
ws.append(["张三", "13800000001", "星辰科技", date(2026, 9, 19), "回电很积极"])
ws.append(["李四", 13800000002, None, None, None])
ws.append(["王五", "138 0000 0003", "海角贸易", "2026/10/1", '他说"再看看"'])
ws.append([None, None, None, None, None])
wb.save(os.path.join(out, "名单.xlsx"))

# 2) 删过工作表的簿子：第一张标签页对应的文件叫 sheet2.xml。
#    openpyxl 保存时会重新编号，所以改名这一步只能手工改 zip。
tmp = os.path.join(out, "_tmp.xlsx")
wb2 = Workbook()
a = wb2.active
a.title = "先删掉的"
b = wb2.create_sheet("真正的名单")
b.append(["姓名", "手机号"])
b.append(["赵六", "13800000006"])
wb2.remove(a)
wb2.save(tmp)

zin = zipfile.ZipFile(tmp)
items = {n: zin.read(n) for n in zin.namelist()}
zin.close()
items["xl/worksheets/sheet2.xml"] = items.pop("xl/worksheets/sheet1.xml")
items["xl/_rels/workbook.xml.rels"] = (
    items["xl/_rels/workbook.xml.rels"].decode().replace("worksheets/sheet1.xml", "worksheets/sheet2.xml").encode()
)
items["[Content_Types].xml"] = (
    items["[Content_Types].xml"].decode().replace("/xl/worksheets/sheet1.xml", "/xl/worksheets/sheet2.xml").encode()
)
with zipfile.ZipFile(os.path.join(out, "删过表.xlsx"), "w", zipfile.ZIP_DEFLATED) as z:
    for n, d in items.items():
        z.writestr(n, d)
os.remove(tmp)
print("写好了：", out)
