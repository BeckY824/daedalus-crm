#!/usr/bin/env python3
"""
企业版导入样例 · 两种平台导出：飞书多维表格、钉钉客户管理。

    pip install openpyxl && python3 scripts/make-ent-import-sample-platforms.py

写到 tests/fixtures/ent-sample-feishu/、ent-sample-dingtalk/（裁判测试 tests/ent-import-platform-judges.test.ts 读它们）。
员工名单用 tests/fixtures/ent-sample/员工与部门.xlsx 那一份（先跑 make-ent-import-sample.py），两边同一批人。
固定随机种子，重跑出来的文件内容一样。数据全是编的。

两种导出的样子是照官方导出格式的「形状」编的，列名、阶段叫法都是常见写法，不是哪一家的真实字段：
  - 飞书多维表格：一个 xlsx 里几张表（客户 / 联系人 / 商机 / 跟进记录），另有一张隐藏表、一张空表；
    「关联」列导出来是对方的名称（没有客户 ID）；人员列一格可以有几个人（「张伟,赵磊」）；
    多选「标签」逗号分隔；复选框导出成 TRUE / FALSE；客户有自动编号「编号」。
  - 钉钉客户管理：客户、商机、跟进记录各一个文件；客户表里一行带一个联系人（联系人姓名 / 手机 / 职位），
    一家有几个联系人就重复几行；有「负责人」「协同人」两列；没有客户编号。

脏数据是故意的，每一种在各自的 说明.md 里列了条数：重名（员工重名、客户重名）、负责人对不上（离职、英文名、重名、
空着、一格几个人）、阶段叫法和我们不一样（有一个规则也认不出，要在对照表里填）、找不到所属客户、缺必填。
"""
from openpyxl import Workbook, load_workbook
from datetime import datetime, timedelta
import os, random, sys

根 = os.path.join(os.path.dirname(__file__), "..", "tests", "fixtures")
rnd = random.Random(20261026)
基准 = datetime(2026, 10, 20, 9, 0)

# ───────────── 员工：和 ent-sample 同一份 ─────────────
员工表 = load_workbook(os.path.join(根, "ent-sample", "员工与部门.xlsx")).active
头 = [c.value for c in next(员工表.iter_rows(max_row=1))]
员工 = [dict(zip(头, [c.value for c in r])) for r in 员工表.iter_rows(min_row=2)]
销售们 = [e for e in 员工 if e["职位"] in ("销售", "销售组长", "大客户经理", "大客户销售") and e["在职状态"] == "在职"]
唯一名销售 = [e for e in 销售们 if sum(1 for x in 员工 if x["姓名"] == e["姓名"]) == 1]

城市 = ["杭州", "宁波", "上海", "苏州", "深圳", "广州", "南京", "合肥"]
字号 = ["曜石", "北辰", "知行", "安澜", "合众", "青松", "朗润", "瀚海", "观澜", "松柏", "尚品", "凌云", "九州", "弘毅", "明德", "清源",
       "丰源", "广厦", "聚贤", "星耀", "睿智", "鸿运", "泰和", "钧天", "沐风", "朝阳", "晨曦", "大道", "久安", "嘉信", "辰光", "蓝海"]
行业 = ["智能装备", "新材料", "冷链物流", "生物科技", "环保科技", "电气", "汽车零部件", "家居"]
姓 = "王李张刘陈杨黄赵吴周徐孙马朱胡郭何高林罗"
名字们 = ["伟", "芳", "娜", "敏", "静", "磊", "军", "洋", "勇", "艳", "杰", "涛", "明", "超", "霞", "平", "刚", "桂英", "建华", "晓东", "海燕", "志强"]
职务 = ["采购经理", "总经理", "技术负责人", "财务", "行政主管", "生产厂长", "采购专员"]
产品 = ["产线改造", "年度框架采购", "仓储系统升级", "设备维保续约", "新厂区配套", "试点项目", "备件采购"]


def 手机(n):
    return f"13{rnd.randint(5, 9)}{n % 10**8:08d}"


def 客户名们(k, 前缀):
    out, 见 = [], set()
    while len(out) < k:
        名 = f"{rnd.choice(城市)}{rnd.choice(字号)}{rnd.choice(行业)}有限公司"
        if 名 not in 见:
            见.add(名)
            out.append(名)
    return out


def 写说明(目录, 标题, 文件说明, 脏):
    with open(os.path.join(目录, "说明.md"), "w", encoding="utf-8") as f:
        f.write(f"# {标题}\n\n由 scripts/make-ent-import-sample-platforms.py 生成，数据全是编的。员工名单和 ../ent-sample/员工与部门.xlsx 同一份。\n\n")
        f.write("## 文件\n\n" + "\n".join(f"- {x}" for x in 文件说明) + "\n\n## 故意放的脏数据\n\n")
        f.write("| 哪一种 | 条数 |\n|---|---|\n" + "\n".join(f"| {k} | {v} |" for k, v in 脏.items()) + "\n")


# ═════════════ 飞书多维表格 ═════════════
目录 = os.path.join(根, "ent-sample-feishu")
os.makedirs(目录, exist_ok=True)
脏 = {}

客户 = []  # {编号, 名, 电话, 负责人(写法)}
for i, 名 in enumerate(客户名们(36, "飞书")):
    客户.append({"编号": i + 1, "名": 名, "电话": 手机(30000000 + i * 7919), "负责人": rnd.choice(销售们)["姓名"]})
# 客户重名：两家同名、不同电话（自动编号不同）——按名称关联的联系人 / 商机 / 跟进认不出是哪家
重名客户 = "南京同名测试科技有限公司"
for k in range(2):
    客户.append({"编号": len(客户) + 1, "名": 重名客户, "电话": 手机(39000000 + k), "负责人": rnd.choice(唯一名销售)["姓名"]})
脏["客户重名（同名不同电话，编号不同）"] = 2

# 负责人：唯一名字认得出；一格几个人取第一个；重名只写名字、离职、英文名、空着 → 落默认负责人
可挑 = [c for c in 客户 if c["名"] != 重名客户]
挑 = rnd.sample(可挑, 12)
for c in 挑[:3]:
    a, b = rnd.sample(唯一名销售, 2)
    c["负责人"] = f"{a['姓名']},{b['姓名']}"
for c, 写法 in zip(挑[3:], ["王芳", "王芳", "李强", "钱进", "钱进", "Kevin Zhang", "", ""]):
    c["负责人"] = 写法
for c in 挑[11:]:
    c["负责人"] = "wangfang.sz@sample.invalid"  # 重名的写邮箱认得出
for c in 可挑:
    if c not in 挑 and c["负责人"] in ("王芳", "李强"):
        c["负责人"] = rnd.choice(唯一名销售)["姓名"]
脏["客户负责人一格几个人（取第一个，提醒）"] = 3
脏["客户负责人只写重名的名字（王芳 / 李强）"] = 3
脏["客户负责人：离职的钱进"] = 2
脏["客户负责人：英文名"] = 1
脏["客户负责人：空着"] = 2

标签们 = ["重点", "华东", "华南", "制造业", "老客户", "展会"]
客户行 = []
for c in 客户:
    客户行.append([
        c["编号"], c["名"], c["负责人"], ",".join(rnd.sample(标签们, rnd.randint(0, 2))), rnd.random() < 0.3,
        rnd.choice(["跟进中", "意向客户", "已成交", "新客户", "沉睡"]), c["电话"],
        rnd.choice(["展会", "转介绍", "官网咨询", "电话开发"]), (基准 - timedelta(days=rnd.randint(20, 500))).strftime("%Y/%m/%d %H:%M"),
        "" if rnd.random() < 0.7 else "年底有预算",
    ])
# 缺必填：客户名称空着
客户行.append([len(客户) + 1, "", "张伟", "", False, "新客户", 手机(39500000), "官网咨询", "2026/03/01 10:00", ""])
脏["客户缺名称（跳过）"] = 1
rnd.shuffle(客户行)

wb = Workbook()
ws = wb.active
ws.title = "客户"
ws.append(["编号", "客户名称", "客户负责人", "标签", "是否重点客户", "客户状态", "联系电话", "客户来源", "创建时间", "备注"])
for r in 客户行:
    ws.append(r)

# 联系人：「所属客户」是关联列，导出来是客户名称
正常 = [c for c in 客户 if c["名"] != 重名客户]
联系行 = []
for c in 正常:
    for _ in range(rnd.choice([1, 1, 2])):
        联系行.append([rnd.choice(姓) + rnd.choice(名字们), c["名"], rnd.choice(职务), 手机(40000000 + len(联系行) * 13),
                     rnd.random() < 0.3, f"fs{len(联系行)}@sample.invalid" if rnd.random() < 0.4 else ""])
for k in range(2):
    联系行.append([rnd.choice(姓) + rnd.choice(名字们), 重名客户, "采购经理", 手机(41000000 + k), False, ""])
脏["联系人关联到重名客户（客户不唯一，跳过）"] = 2
联系行.append([rnd.choice(姓) + rnd.choice(名字们), "查无此司有限公司", "总经理", 手机(41100000), False, ""])
脏["联系人关联的客户不存在（找不到客户，跳过）"] = 1
rnd.shuffle(联系行)
ws = wb.create_sheet("联系人")
ws.append(["联系人姓名", "所属客户", "职务", "手机号", "关键联系人", "邮箱"])
for r in 联系行:
    ws.append(r)

# 商机：阶段叫法和我们不一样；「搁置」规则认不出，要在对照表里填
飞书阶段 = ["初步接洽", "需求确认", "方案报价", "商务谈判", "赢单", "输单", "搁置"]
商机行 = []
for c in rnd.sample(正常, 24):
    d = 基准 + timedelta(days=rnd.randint(-60, 120))
    商机行.append([f"{c['名'][:4]}{rnd.choice(产品)}", c["名"], rnd.choices(飞书阶段, [5, 5, 5, 3, 3, 2, 2])[0],
                 f"¥{rnd.choice([3, 8, 15, 30, 60]) * 10000:,}.00", d.strftime("%Y/%m/%d"), c["负责人"] if rnd.random() < 0.5 else rnd.choice(唯一名销售)["姓名"]])
# 保证「搁置」至少两条（对照表要填它）
for r in 商机行[:2]:
    r[2] = "搁置"
商机行.append(["孤儿商机", "查无此司有限公司", "初步接洽", "¥10,000.00", "2026/12/01", "张伟"])
脏["商机阶段「搁置」规则认不出（对照表里填）"] = sum(1 for r in 商机行 if r[2] == "搁置")
脏["商机关联的客户不存在（找不到客户，跳过）"] = 1
rnd.shuffle(商机行)
ws = wb.create_sheet("商机")
ws.append(["商机名称", "关联客户", "销售阶段", "预计金额", "预计成交日期", "负责人"])
for r in 商机行:
    ws.append(r)

# 跟进记录
飞书方式 = ["电话", "上门拜访", "微信", "视频会议", "邮件", "其他"]
跟进行 = []
for k in range(60):
    c = rnd.choice(正常)
    t = 基准 - timedelta(days=rnd.randint(0, 200), minutes=rnd.randint(0, 600))
    跟进行.append([c["名"], rnd.choice(飞书方式), f"第{k + 1}次跟进：{rnd.choice(['聊了预算', '约了下周演示', '对方在比价', '寄了样品', '催回款'])}",
                 t.strftime("%Y/%m/%d %H:%M"), rnd.choice(唯一名销售)["姓名"]])
跟进行.append([重名客户, "电话", "重名客户的一次跟进", "2026/10/01 10:00", "张伟"])
脏["跟进关联到重名客户（客户不唯一，跳过）"] = 1
rnd.shuffle(跟进行)
ws = wb.create_sheet("跟进记录")
ws.append(["关联客户", "跟进方式", "跟进内容", "跟进时间", "跟进人"])
for r in 跟进行:
    ws.append(r)

# 隐藏表、空表：不该读
ws = wb.create_sheet("字段说明")
ws.append(["字段", "说明"])
ws.append(["销售阶段", "按公司流程填"])
ws.sheet_state = "hidden"
wb.create_sheet("仪表盘")
wb.save(os.path.join(目录, "客户管理.xlsx"))

写说明(目录, "飞书多维表格导出样例", [
    "客户管理.xlsx：客户 / 联系人 / 商机 / 跟进记录四张表；「字段说明」是隐藏表、「仪表盘」是空表，都不该读",
    "客户有自动编号「编号」；联系人、商机、跟进按客户名称关联（没有客户 ID）",
    "人员列一格可能几个人（逗号分隔）；标签是多选、逗号分隔；复选框是 TRUE / FALSE",
], 脏)

# ═════════════ 钉钉客户管理 ═════════════
目录 = os.path.join(根, "ent-sample-dingtalk")
os.makedirs(目录, exist_ok=True)
脏 = {}
rnd = random.Random(20261027)

客户 = []
for i, 名 in enumerate(客户名们(34, "钉钉")):
    客户.append({"名": 名, "电话": f"0571-{8000_0000 + i * 37:08d}", "负责人": rnd.choice(销售们)["姓名"], "协同人": rnd.choice(唯一名销售)["姓名"] if rnd.random() < 0.4 else ""})
for k, 写法 in enumerate(["李强", "李强", "钱进", "Kevin", "", "liqiang.sh@sample.invalid"]):
    客户[k]["负责人"] = 写法
for c in 客户[len(["李强", "李强", "钱进", "Kevin", "", "liqiang.sh@sample.invalid"]):]:
    if c["负责人"] in ("王芳", "李强"):
        c["负责人"] = rnd.choice(唯一名销售)["姓名"]
脏["客户负责人只写重名的名字（李强）"] = 2
脏["客户负责人：离职的钱进"] = 1
脏["客户负责人：英文名"] = 1
脏["客户负责人：空着"] = 1
# 同名客户：两家同名、不同电话（没有编号，按名称 + 电话区分）
重名客户 = "苏州同名测试精密有限公司"
for k in range(2):
    客户.append({"名": 重名客户, "电话": f"0512-{6600_0000 + k:08d}", "负责人": rnd.choice(唯一名销售)["姓名"], "协同人": ""})
脏["客户重名（同名不同电话）"] = 2

钉钉状态 = ["潜在客户", "意向客户", "成交客户", "流失客户"]
客户行 = []
for c in 客户:
    共 = [c["名"], c["电话"], rnd.choice(钉钉状态), rnd.choice(["地推", "转介绍", "线上"]), rnd.choice(["A", "B", "C"]), c["负责人"], c["协同人"],
         f"{rnd.choice(城市)}市高新区{rnd.randint(1, 999)}号", (基准 - timedelta(days=rnd.randint(30, 400))).strftime("%Y-%m-%d %H:%M:%S")]
    人数 = 0 if c["名"] == 重名客户 else rnd.choice([0, 1, 1, 2, 3])
    if 人数 == 0:
        客户行.append(共 + ["", "", ""])
    for _ in range(人数):
        客户行.append(共 + [rnd.choice(姓) + rnd.choice(名字们), 手机(50000000 + len(客户行) * 17), rnd.choice(职务)])
脏["一家几个联系人，客户行重复（拆联系人后客户只留一行）"] = len(客户行) - len(客户)
客户行.append(["", "0571-80009999", "潜在客户", "线上", "C", "张伟", "", "", "2026-03-01 10:00:00", "", "", ""])
脏["客户缺名称（跳过）"] = 1
rnd.shuffle(客户行)
wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.append(["客户名称", "电话", "客户状态", "客户来源", "客户级别", "负责人", "协同人", "地址", "创建时间", "联系人姓名", "联系人手机", "联系人职位"])
for r in 客户行:
    ws.append(r)
wb.save(os.path.join(目录, "客户.xlsx"))

# 商机：阶段叫法和我们不一样（规则大都认得出）
钉钉阶段 = ["初步沟通", "需求分析", "方案/报价", "谈判审核", "赢单", "输单"]
正常 = [c for c in 客户 if c["名"] != 重名客户]
商机行 = []
for c in rnd.sample(正常, 22):
    d = 基准 + timedelta(days=rnd.randint(-30, 150))
    商机行.append([f"{c['名'][2:6]}{rnd.choice(产品)}", c["名"], rnd.choice(钉钉阶段), rnd.choice([30000, 80000, 150000, 600000]), d.strftime("%Y-%m-%d"),
                 c["负责人"] if c["负责人"] and rnd.random() < 0.5 else rnd.choice(唯一名销售)["姓名"]])
商机行.append(["重名客户的商机", 重名客户, "初步沟通", 50000, "2026-12-31", "张伟"])
脏["商机关联到重名客户（客户不唯一，跳过）"] = 1
rnd.shuffle(商机行)
wb = Workbook()
ws = wb.active
ws.append(["商机名称", "客户名称", "商机阶段", "预计成交金额", "预计成交日期", "负责人"])
for r in 商机行:
    ws.append(r)
wb.save(os.path.join(目录, "商机.xlsx"))

# 跟进记录
钉钉方式 = ["电话沟通", "上门拜访", "微信沟通", "其他"]
跟进行 = []
for k in range(50):
    c = rnd.choice(正常)
    t = 基准 - timedelta(days=rnd.randint(0, 150), minutes=rnd.randint(0, 600))
    跟进行.append([c["名"], rnd.choice(钉钉方式), f"钉钉跟进{k + 1}：{rnd.choice(['约了拜访', '发了报价单', '对方要样品', '在走审批'])}", t.strftime("%Y-%m-%d %H:%M:%S"),
                 rnd.choice(唯一名销售)["姓名"]])
跟进行.append(["查无此司有限公司", "电话沟通", "找不到客户的一次跟进", "2026-10-01 10:00:00", "张伟"])
脏["跟进的客户不存在（找不到客户，跳过）"] = 1
rnd.shuffle(跟进行)
wb = Workbook()
ws = wb.active
ws.append(["客户名称", "跟进方式", "跟进内容", "跟进时间", "跟进人"])
for r in 跟进行:
    ws.append(r)
wb.save(os.path.join(目录, "跟进记录.xlsx"))

写说明(目录, "钉钉客户管理导出样例", [
    "客户.xlsx：一行一家客户带一个联系人（联系人姓名 / 手机 / 职位），几个联系人就重复几行；有负责人、协同人两列；没有客户编号",
    "商机.xlsx、跟进记录.xlsx：按客户名称关联",
], 脏)
print("ok")
